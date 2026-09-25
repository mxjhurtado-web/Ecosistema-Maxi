#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
TEMIS Package Builder Service
Generates:
1. Formatted Word Bitácora (.docx) with embedded screenshot images and SIPOC table
2. TEMIS Web Project Package (.temis.json) for 1-click cloud sync
3. Subtitle / Transcript files (.vtt / .txt)
"""

import os
import json
import datetime
import logging
from typing import List, Dict, Any, Optional

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml, OxmlElement
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    Document = None

logger = logging.getLogger("temis_media_studio")


class TemisPackageBuilder:
    """Builds Word documents with embedded images and TEMIS web packages"""

    @staticmethod
    def _set_cell_bg(cell, hex_color: str):
        """Set background color of a docx table cell"""
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    @classmethod
    def build_word_bitacora(
        cls,
        analysis_data: Dict[str, Any],
        transcript_segments: List[Dict[str, Any]],
        keyframes: List[Dict[str, Any]],
        output_docx_path: str,
        video_filename: str
    ) -> str:
        """
        Generate executive Word Bitácora document with embedded screenshot keyframes
        """
        if Document is None:
            raise ImportError("python-docx is not installed")

        doc = Document()
        charter = analysis_data.get("project_charter", {})
        steps = analysis_data.get("process_steps", [])
        sipoc = analysis_data.get("sipoc", [])

        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Document Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run("TEMIS MEDIA STUDIO — BITÁCORA DE PROCESO")
        title_run.font.name = "Arial"
        title_run.font.size = Pt(16)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x17, 0x28, 0x3C)

        subtitle_p = doc.add_paragraph()
        subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = subtitle_p.add_run(charter.get("project_name", "Levantamiento de Proceso"))
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(13)
        sub_run.font.bold = True
        sub_run.font.color.rgb = RGBColor(0x1E, 0x5A, 0x9A)

        # Metadata Bar
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
        meta_run = meta_p.add_run(f"Código: {charter.get('project_code', 'PRJ-001')}  |  Archivo: {video_filename}  |  Fecha: {now_str}")
        meta_run.font.size = Pt(9)
        meta_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 1. FICHA DEL PROCESO
        h1 = doc.add_heading("1. Ficha del Proceso y Resumen Ejecutivo", level=1)
        h1.paragraph_format.space_before = Pt(10)
        h1.paragraph_format.space_after = Pt(4)

        charter_table = doc.add_table(rows=4, cols=2)
        charter_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        charter_table.autofit = False

        col_widths = [Inches(1.8), Inches(4.8)]
        charter_data = [
            ("Propósito:", charter.get("purpose", "Estandarización de proceso")),
            ("Alcance:", charter.get("scope", "Operación general")),
            ("Sistemas Clave:", charter.get("target_system", "Sistemas Corporativos")),
            ("Área / Sponsor:", charter.get("sponsor", "Operaciones"))
        ]

        for r_idx, (label, val) in enumerate(charter_data):
            row = charter_table.rows[r_idx]
            cell_lbl = row.cells[0]
            cell_val = row.cells[1]
            
            cell_lbl.width = col_widths[0]
            cell_val.width = col_widths[1]
            cls._set_cell_bg(cell_lbl, "F1F5F9")
            
            p_lbl = cell_lbl.paragraphs[0]
            p_lbl.add_run(label).bold = True
            p_lbl.runs[0].font.size = Pt(9.5)
            
            p_val = cell_val.paragraphs[0]
            p_val.add_run(str(val))
            p_val.runs[0].font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Executive Summary Paragraph
        if charter.get("executive_summary"):
            exec_p = doc.add_paragraph()
            exec_run = exec_p.add_run(charter.get("executive_summary"))
            exec_run.font.size = Pt(10)
            exec_p.paragraph_format.space_after = Pt(10)

        # 2. MATRIZ SIPOC
        if sipoc:
            h2 = doc.add_heading("2. Matriz SIPOC de Alto Nivel", level=1)
            h2.paragraph_format.space_before = Pt(12)
            h2.paragraph_format.space_after = Pt(4)

            sipoc_table = doc.add_table(rows=len(sipoc) + 1, cols=6)
            sipoc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            headers = ["ID", "Proveedor (S)", "Entrada (I)", "Proceso (P)", "Salida (O)", "Cliente (C)"]
            hdr_row = sipoc_table.rows[0]
            for c_idx, h_text in enumerate(headers):
                cell = hdr_row.cells[c_idx]
                cls._set_cell_bg(cell, "1E5A9A")
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(h_text)
                r.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            for r_idx, row_data in enumerate(sipoc):
                row = sipoc_table.rows[r_idx + 1]
                vals = [
                    row_data.get("id", f"{r_idx+1}.0"),
                    row_data.get("supplier", ""),
                    row_data.get("input", ""),
                    row_data.get("process", ""),
                    row_data.get("output", ""),
                    row_data.get("customer", "")
                ]
                bg_color = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
                for c_idx, v in enumerate(vals):
                    cell = row.cells[c_idx]
                    cls._set_cell_bg(cell, bg_color)
                    p = cell.paragraphs[0]
                    r = p.add_run(str(v))
                    r.font.size = Pt(8.5)

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 3. BITÁCORA DETALLADA PASO A PASO CON CAPTURAS
        h3 = doc.add_heading("3. Bitácora de Ejecución con Evidencia Visual", level=1)
        h3.paragraph_format.space_before = Pt(12)
        h3.paragraph_format.space_after = Pt(6)

        # Create quick map of keyframes by filename
        kf_map = {k.get("filename"): k.get("path") for k in keyframes if k.get("path") and os.path.exists(k.get("path"))}

        for step in steps:
            step_num = step.get("step_number", 1)
            step_title = step.get("title", f"Paso {step_num}")
            step_time = step.get("timestamp", "00:00:00")
            step_actor = step.get("actor", "Operador")
            step_sys = step.get("system", "Sistema")
            step_desc = step.get("description", "")
            attached_img = step.get("attached_screenshot", "")
            img_caption = step.get("screenshot_caption", f"Captura del paso {step_num}")

            # Step Header
            step_p = doc.add_paragraph()
            step_p.paragraph_format.space_before = Pt(8)
            step_p.paragraph_format.space_after = Pt(2)
            
            r_num = step_p.add_run(f"Paso {step_num}: {step_title} ")
            r_num.bold = True
            r_num.font.size = Pt(11)
            r_num.font.color.rgb = RGBColor(0x17, 0x28, 0x3C)

            r_time = step_p.add_run(f"[{step_time}]")
            r_time.font.size = Pt(9.5)
            r_time.font.bold = True
            r_time.font.color.rgb = RGBColor(0x7C, 0x3A, 0xED)

            # Metadata tags
            tag_p = doc.add_paragraph()
            tag_p.paragraph_format.space_after = Pt(3)
            r_tag = tag_p.add_run(f"👤 Responsable: {step_actor}  |  💻 Sistema: {step_sys}")
            r_tag.font.size = Pt(8.5)
            r_tag.font.italic = True
            r_tag.font.color.rgb = RGBColor(0x52, 0x65, 0x7A)

            # Description
            desc_p = doc.add_paragraph()
            desc_p.paragraph_format.space_after = Pt(6)
            r_desc = desc_p.add_run(step_desc)
            r_desc.font.size = Pt(9.5)

            # Insert Image if attached
            img_path = kf_map.get(attached_img)
            if not img_path and keyframes:
                # If exact name not found, attach next available
                kf_idx = min(step_num - 1, len(keyframes) - 1)
                img_path = keyframes[kf_idx].get("path")

            if img_path and os.path.exists(img_path):
                try:
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(4)
                    img_p.paragraph_format.space_after = Pt(2)
                    doc_run = img_p.add_run()
                    doc_run.add_picture(img_path, width=Inches(5.6))
                    
                    # Caption
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(8)
                    cap_run = cap_p.add_run(f"Figura {step_num}: {img_caption}")
                    cap_run.font.size = Pt(8)
                    cap_run.font.italic = True
                    cap_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
                except Exception as img_err:
                    logger.warning(f"Could not embed picture {img_path}: {img_err}")

        # 4. APÉNDICE: TRANSCRIPCIÓN COMPLETA
        if transcript_segments:
            h4 = doc.add_heading("4. Apéndice: Transcripción Literal de la Grabación", level=1)
            h4.paragraph_format.space_before = Pt(14)
            h4.paragraph_format.space_after = Pt(4)

            for seg in transcript_segments:
                seg_p = doc.add_paragraph()
                seg_p.paragraph_format.space_after = Pt(2)
                
                r_ts = seg_p.add_run(f"[{seg.get('timestamp_start', '00:00')}] ")
                r_ts.font.size = Pt(8.5)
                r_ts.font.bold = True
                r_ts.font.color.rgb = RGBColor(0x1E, 0x5A, 0x9A)

                r_txt = seg_p.add_run(seg.get("text", ""))
                r_txt.font.size = Pt(8.5)

        # Save document
        doc.save(output_docx_path)
        logger.info(f"Word Bitácora saved successfully to: {output_docx_path}")
        return output_docx_path

    @classmethod
    def build_temis_json_package(
        cls,
        analysis_data: Dict[str, Any],
        transcript_segments: List[Dict[str, Any]],
        keyframes: List[Dict[str, Any]],
        output_json_path: str,
        video_filename: str
    ) -> str:
        """
        Generate .temis.json project package compatible with TEMIS Web 1-click import
        """
        charter = analysis_data.get("project_charter", {})
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        package = {
            "schema_version": "2.0",
            "generator": "TEMIS Media Studio Desktop",
            "generated_at": now_str,
            "project_id": charter.get("project_code", "PRJ-001"),
            "project_name": charter.get("project_name", "Nuevo Proceso"),
            "project_code": charter.get("project_code", "PRJ-001"),
            "source_video_file": video_filename,
            "charter": {
                "project_name": charter.get("project_name", "Nuevo Proceso"),
                "project_code": charter.get("project_code", "PRJ-001"),
                "purpose": charter.get("purpose", ""),
                "scope": charter.get("scope", ""),
                "target_system": charter.get("target_system", ""),
                "sponsor": charter.get("sponsor", "Operaciones"),
                "executive_summary": charter.get("executive_summary", "")
            },
            "sipoc_rows": analysis_data.get("sipoc", []),
            "diagram": {
                "title": charter.get("project_name", "Diagrama de Proceso"),
                "nodes": analysis_data.get("bpmn_nodes", []),
                "edges": analysis_data.get("bpmn_edges", [])
            },
            "process_steps": analysis_data.get("process_steps", []),
            "transcript_segments": transcript_segments,
            "extracted_keyframes": [
                {
                    "index": k.get("index"),
                    "filename": k.get("filename"),
                    "timestamp_formatted": k.get("timestamp_formatted")
                }
                for k in keyframes
            ]
        }

        with open(output_json_path, "w", encoding="utf-8") as f:
            json.dump(package, f, indent=2, ensure_ascii=False)

        logger.info(f"TEMIS JSON package saved to: {output_json_path}")
        return output_json_path

    @classmethod
    def export_vtt(cls, segments: List[Dict[str, Any]], output_vtt_path: str) -> str:
        """Export subtitle VTT file"""
        lines = ["WEBVTT", ""]
        for s in segments:
            idx = s.get("index", 1)
            t_start = s.get("timestamp_start", "00:00:00")
            t_end = s.get("timestamp_end", "00:00:10")
            text = s.get("text", "")
            # Ensure format 00:00:00.000
            if len(t_start.split(":")) == 2: t_start = f"00:{t_start}.000"
            elif "." not in t_start: t_start = f"{t_start}.000"
            if len(t_end.split(":")) == 2: t_end = f"00:{t_end}.000"
            elif "." not in t_end: t_end = f"{t_end}.000"

            lines.append(f"{idx}")
            lines.append(f"{t_start} --> {t_end}")
            lines.append(f"{text}")
            lines.append("")

        with open(output_vtt_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return output_vtt_path

    @classmethod
    def export_txt(cls, segments: List[Dict[str, Any]], output_txt_path: str, title: str) -> str:
        """Export plain text transcript with timestamps"""
        lines = [
            f"============================================================",
            f"TEMIS MEDIA STUDIO — TRANSCRIPCIÓN CON MINUTAJES",
            f"Proyecto: {title}",
            f"Fecha: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}",
            f"============================================================",
            ""
        ]
        for s in segments:
            lines.append(f"[{s.get('timestamp_start', '00:00')}] {s.get('speaker', 'Participante')}:")
            lines.append(f"  {s.get('text', '')}")
            lines.append("")

        with open(output_txt_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return output_txt_path

    @classmethod
    def create_zip_package(cls, export_subfolder: str, base_name: str) -> str:
        """
        Create a compact .zip bundle containing Word bitácora, .temis.json, 
        transcript .vtt/.txt, and capturas/ folder, explicitly excluding raw audio.wav 
        or heavy video files so the upload to TEMIS Web is ~1-2 MB.
        """
        import zipfile
        zip_filename = f"Paquete_TEMIS_{base_name}.zip"
        zip_path = os.path.join(export_subfolder, zip_filename)
        
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
            for root, dirs, files in os.walk(export_subfolder):
                for file in files:
                    file_path = os.path.join(root, file)
                    rel_path = os.path.relpath(file_path, export_subfolder)
                    # Exclude raw audio, video or other zip files
                    if file.lower().endswith((".wav", ".mp4", ".mov", ".mkv", ".webm", ".avi", ".zip")):
                        continue
                    z.write(file_path, arcname=rel_path)
                    
        logger.info(f"Compact TEMIS ZIP package created at: {zip_path} ({os.path.getsize(zip_path)/1024:.1f} KB)")
        return zip_path

