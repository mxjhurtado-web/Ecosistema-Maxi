#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
SIPOC Excel Exporter for TEMIS
Generates styled professional Microsoft Excel (.xlsx) workbooks 
matching the official Six Sigma SIPOC template.
"""

import io
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_sipoc_to_excel(
    project_name: str,
    project_purpose: str,
    sipoc_rows: List[Dict[str, Any]],
    customer_requirements: str = ""
) -> io.BytesIO:
    """Generate professional styled Six Sigma SIPOC Excel workbook in memory"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matriz SIPOC"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    COLOR_PRIMARY = "1E3A8A"      # Dark Blue
    COLOR_SUPPLIER = "1E40AF"     # Royal Blue
    COLOR_INPUT = "0284C7"        # Sky Blue
    COLOR_PROCESS = "0F766E"      # Teal
    COLOR_OUTPUT = "C2410C"       # Orange
    COLOR_CUSTOMER = "6D28D9"     # Purple
    COLOR_REQ = "B45309"          # Amber
    COLOR_ZEBRA = "F8FAFC"        # Light Slate

    # Borders
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    header_border = Border(
        left=Side(style="medium", color="FFFFFF"),
        right=Side(style="medium", color="FFFFFF"),
        top=Side(style="medium", color="1E3A8A"),
        bottom=Side(style="medium", color="1E3A8A")
    )

    # 1. Main Title Banner (Row 2)
    ws.merge_cells("B2:G2")
    cell_title = ws["B2"]
    cell_title.value = f"PLANTILLA DE MATRIZ SIPOC — {project_name.upper()}"
    cell_title.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    cell_title.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    cell_title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 32

    # 2. Purpose Subtitle (Row 3)
    ws.merge_cells("B3:G3")
    cell_sub = ws["B3"]
    cell_sub.value = f"Propósito: {project_purpose}" if project_purpose else "Herramienta Six Sigma para mapeo de alto nivel de procesos de inicio a fin."
    cell_sub.font = Font(name="Calibri", size=10, italic=True, color="475569")
    cell_sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 20

    # 3. SIPOC Initial Letters (Row 5)
    sipoc_headers = [
        ("B", "S", "S U P P L I E R S", COLOR_SUPPLIER),
        ("C", "I", "I N P U T", COLOR_INPUT),
        ("D", "P", "P R O C E S S", COLOR_PROCESS),
        ("E", "O", "O U T P U T", COLOR_OUTPUT),
        ("F", "C", "C U S T O M E R", COLOR_CUSTOMER),
        ("G", "R", "R E Q U I R E M E N T S", COLOR_REQ)
    ]

    ws.row_dimensions[5].height = 24
    for col_letter, letter, title, color in sipoc_headers:
        c = ws[f"{col_letter}5"]
        c.value = f"{letter} - {title}"
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = header_border

    # 4. SIPOC Column Explanations (Row 6)
    explanations = [
        ("B", "Quién suministra lo que se requiere para ejecutar el proceso (Entrada)."),
        ("C", "Recurso / insumo proporcionado por el proveedor para el proceso."),
        ("D", "Actividades secuenciales realizadas para convertir entradas en salidas."),
        ("E", "Recurso o entregable resultante de la actividad realizada."),
        ("F", "Receptor o beneficiario de la salida creada."),
        ("G", "Criterios de calidad, tiempo o especificaciones requeridas.")
    ]
    ws.row_dimensions[6].height = 36
    for col_letter, exp in explanations:
        c = ws[f"{col_letter}6"]
        c.value = exp
        c.font = Font(name="Calibri", size=8.5, italic=True, color="64748B")
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    # 5. Column Headers (Row 7)
    labels = [
        ("B", "PROVEEDORES"),
        ("C", "ENTRADA"),
        ("D", "PROCESO (Paso 1.0..N)"),
        ("E", "SALIDA"),
        ("F", "CLIENTE"),
        ("G", "REQUISITOS")
    ]
    ws.row_dimensions[7].height = 22
    for col_letter, lbl in labels:
        c = ws[f"{col_letter}7"]
        c.value = lbl
        c.font = Font(name="Calibri", size=9.5, bold=True, color="334155")
        c.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # 6. Data Rows
    current_row = 8
    
    # Ensure minimum 10 rows
    rows_to_render = list(sipoc_rows)
    if len(rows_to_render) < 10:
        for i in range(len(rows_to_render) + 1, 11):
            rows_to_render.append({
                "id": str(i),
                "step_num": f"{i}.0",
                "provider": "",
                "input": "",
                "step": f"Paso {i}.0",
                "output": "",
                "customer": "",
                "requirements": ""
            })

    for idx, r in enumerate(rows_to_render):
        ws.row_dimensions[current_row].height = 24
        fill_color = "FFFFFF" if idx % 2 == 0 else COLOR_ZEBRA
        
        step_val = r.get("step") or r.get("process") or f"{idx+1}.0"
        if not step_val.startswith(f"{idx+1}.0") and not any(step_val.startswith(f"{n}.") for n in range(1, 20)):
            step_val = f"{idx+1}.0 {step_val}"

        row_data = [
            ("B", r.get("provider", "") or r.get("supplier", "")),
            ("C", r.get("input", "")),
            ("D", step_val),
            ("E", r.get("output", "")),
            ("F", r.get("customer", "")),
            ("G", r.get("requirements", "") or customer_requirements)
        ]

        for col_letter, val in row_data:
            c = ws[f"{col_letter}{current_row}"]
            c.value = val
            c.font = Font(name="Calibri", size=9.5, color="1E293B", bold=(col_letter == "D"))
            c.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            c.alignment = Alignment(horizontal="left" if col_letter in ["B", "C", "D", "E", "F", "G"] else "center", vertical="center")
            c.border = thin_border

        current_row += 1

    # 7. Bottom Customer Requirements Summary
    current_row += 1
    ws.merge_cells(f"B{current_row}:G{current_row}")
    cell_req = ws[f"B{current_row}"]
    cell_req.value = f"REQUISITOS DEL CLIENTE / NOTAS DE CALIDAD: {customer_requirements}" if customer_requirements else "REQUISITOS DEL CLIENTE: Cumplimiento de tiempos de respuesta (SLA), integridad de datos en Chronos y confirmación de satisfacción."
    cell_req.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
    cell_req.fill = PatternFill(start_color=COLOR_REQ, end_color=COLOR_REQ, fill_type="solid")
    cell_req.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[current_row].height = 28

    # Column Widths
    col_widths = {
        "A": 3,
        "B": 24,  # Suppliers
        "C": 26,  # Input
        "D": 34,  # Process
        "E": 26,  # Output
        "F": 24,  # Customer
        "G": 28   # Requirements
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
