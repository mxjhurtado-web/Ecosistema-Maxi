#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Transcript Exporter Service for TEMIS
Generates official audit backup files in Word (.docx) and plain text (.txt)
for audio/video interview recordings and meeting minutes.
"""

import io
import datetime
from typing import List, Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.models.narrative_source_model import ParagraphBlock


def set_cell_background(cell, fill_hex: str):
    """Set background color for table cell in DOCX"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex.replace("#", ""))
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in DXA"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


class TranscriptExporter:
    """Service to export audit-ready interview transcripts in DOCX and TXT formats"""

    @classmethod
    def export_to_txt(
        cls,
        project_name: str,
        source_filename: str,
        duration_str: str,
        uploaded_by: str,
        blocks: List[Dict[str, Any]]
    ) -> str:
        """Generate structured text transcript for backup"""
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        lines = [
            "=" * 80,
            f"MINUTA & TRANSCRIPCIÓN OFICIAL DE LEVANTAMIENTO DE PROCESO",
            f"PROYECTO: {project_name.upper()}",
            "=" * 80,
            f"Archivo Fuente:   {source_filename}",
            f"Duración:         {duration_str}",
            f"Fecha de Carga:   {now_str}",
            f"Responsable:      {uploaded_by}",
            f"Generado por:     TEMIS Web Flow (Ecosistema Orbi / Maxi)",
            "=" * 80,
            "",
            "--- REGISTRO DE DIÁLOGO Y MINUTAJE ---",
            ""
        ]

        for b in blocks:
            ts_start = b.get("timestamp_start") or "00:00:00"
            ts_end = b.get("timestamp_end") or ""
            ts_str = f"[{ts_start} - {ts_end}]" if ts_end else f"[{ts_start}]"
            speaker = b.get("speaker") or "Participante"
            text = b.get("text") or ""
            lines.append(f"{ts_str} [{speaker}]:")
            lines.append(f"  {text}")
            lines.append("")

        return "\n".join(lines)

    @classmethod
    def export_to_docx(
        cls,
        project_name: str,
        source_filename: str,
        duration_str: str,
        uploaded_by: str,
        blocks: List[Dict[str, Any]]
    ) -> io.BytesIO:
        """Generate formatted Word .docx document for compliance and project backup"""
        doc = docx.Document()

        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title
        title_p = doc.add_paragraph()
        run_title = title_p.add_run("MINUTA Y TRANSCRIPCIÓN DE ENTREVISTA")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(14)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(23, 40, 60)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub_p = doc.add_paragraph()
        run_sub = sub_p.add_run(f"Proyecto: {project_name}")
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(82, 101, 122)
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # Metadata Summary Table
        meta_table = doc.add_table(rows=4, cols=2)
        meta_table.autofit = False

        headers_data = [
            ("Archivo Fuente Multimedia:", source_filename),
            ("Duración del Registro:", duration_str),
            ("Fecha y Hora de Carga:", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Analista / Responsable:", uploaded_by)
        ]

        for idx, (label, val) in enumerate(headers_data):
            row = meta_table.rows[idx]
            cell_lbl = row.cells[0]
            cell_val = row.cells[1]

            cell_lbl.width = Inches(2.5)
            cell_val.width = Inches(4.5)

            set_cell_background(cell_lbl, "F1F5F9")
            set_cell_margins(cell_lbl, 80, 80, 120, 120)
            set_cell_margins(cell_val, 80, 80, 120, 120)

            p_lbl = cell_lbl.paragraphs[0]
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.name = "Calibri"
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.bold = True
            r_lbl.font.color.rgb = RGBColor(30, 90, 154)

            p_val = cell_val.paragraphs[0]
            r_val = p_val.add_run(str(val))
            r_val.font.name = "Calibri"
            r_val.font.size = Pt(9.5)
            r_val.font.color.rgb = RGBColor(23, 40, 60)

        doc.add_paragraph()

        # Section Heading
        h_p = doc.add_paragraph()
        run_h = h_p.add_run("Registro Cronológico de Declaraciones")
        run_h.font.name = "Calibri"
        run_h.font.size = Pt(12)
        run_h.font.bold = True
        run_h.font.color.rgb = RGBColor(30, 90, 154)

        # Content blocks
        for b in blocks:
            ts_start = b.get("timestamp_start") or "00:00:00"
            ts_end = b.get("timestamp_end") or ""
            ts_str = f"[{ts_start} - {ts_end}]" if ts_end else f"[{ts_start}]"
            speaker = b.get("speaker") or "Participante"
            text = b.get("text") or ""

            p_entry = doc.add_paragraph()
            p_entry.paragraph_format.space_before = Pt(4)
            p_entry.paragraph_format.space_after = Pt(2)

            r_ts = p_entry.add_run(f"{ts_str} ")
            r_ts.font.name = "Calibri"
            r_ts.font.size = Pt(9)
            r_ts.font.bold = True
            r_ts.font.color.rgb = RGBColor(124, 58, 237)

            r_spk = p_entry.add_run(f"[{speaker}]: ")
            r_spk.font.name = "Calibri"
            r_spk.font.size = Pt(9.5)
            r_spk.font.bold = True
            r_spk.font.color.rgb = RGBColor(30, 90, 154)

            r_txt = p_entry.add_run(text)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = RGBColor(51, 65, 85)

        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)
        return bio
