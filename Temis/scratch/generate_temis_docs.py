#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Script to generate the 3 official corporate documentation files for PROYECTO TEMIS (.docx):
1. Roadmap del Proyecto TEMIS 2026
2. Marco de Trabajo para la Gestión del Proyecto TEMIS
3. Informe Diagnóstico General y Análisis de Procesos - Proyecto TEMIS
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = r"C:\Users\User\Ecosistema-Maxi\Temis\docs"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Color Palette
HEX_PRIMARY = "1E3A8A"     # Deep Blue
HEX_SECONDARY = "2563EB"   # Royal Blue
HEX_TEAL = "0F766E"        # Teal
HEX_LIGHT_BG = "F1F5F9"    # Light Slate
HEX_BORDER = "CBD5E1"      # Border Gray
HEX_TEXT_DARK = "0F172A"   # Dark Navy Text
HEX_TEXT_MUTED = "475569"  # Slate Muted

COLOR_PRIMARY = RGBColor(30, 58, 138)
COLOR_SECONDARY = RGBColor(37, 99, 235)
COLOR_TEAL = RGBColor(15, 118, 110)
COLOR_TEXT_DARK = RGBColor(15, 23, 42)
COLOR_MUTED = RGBColor(71, 85, 105)


def set_cell_background(cell, hex_color):
    """Set background color of a table cell"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa (1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_document_header(doc, title, subtitle, category="DOCUMENTO OFICIAL DE GOBERNANZA"):
    """Add professional title banner"""
    p_cat = doc.add_paragraph()
    r_cat = p_cat.add_run(f"TEMIS PROCESS SUITE  |  {category.upper()}")
    r_cat.font.name = "Calibri"
    r_cat.font.size = Pt(9)
    r_cat.font.bold = True
    r_cat.font.color.rgb = COLOR_SECONDARY
    p_cat.paragraph_format.space_after = Pt(4)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(6)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_MUTED
    p_sub.paragraph_format.space_after = Pt(18)

    # Metadata table
    tbl = doc.add_table(rows=2, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    meta = [
        ("Proyecto:", "TEMIS (Process Suite & Governance)", "Periodo:", "Enero 2026 – Diciembre 2026"),
        ("Versión:", "1.0 Oficial (Web Cloud SaaS)", "Estado:", "Aprobado / Línea Base")
    ]
    for row_idx, data in enumerate(meta):
        row = tbl.rows[row_idx]
        for col_idx in range(4):
            cell = row.cells[col_idx]
            set_cell_background(cell, "F8FAFC" if col_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(data[col_idx])
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if col_idx % 2 == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)


def add_heading_1(doc, text):
    h = doc.add_paragraph()
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    return h


def add_heading_2(doc, text):
    h = doc.add_paragraph()
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_SECONDARY
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    return h


def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_TEXT_DARK
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_TEXT_DARK
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    return p


# ==============================================================================
# 1. ROADMAP PROYECTO TEMIS 2026
# ==============================================================================
def create_roadmap_document():
    doc = docx.Document()
    add_document_header(
        doc,
        title="Roadmap Estratégico del Proyecto TEMIS 2026",
        subtitle="Cronograma integral, 7 fases metodológicas e hitos de evolución tecnológica hacia Web SaaS Cloud.",
        category="PLANIFICACIÓN ESTRATÉGICA & CRONOGRAMA"
    )

    add_heading_1(doc, "1. Visión y Objetivos del Roadmap 2026")
    add_body_p(doc, "El Proyecto TEMIS tiene como propósito consolidar una plataforma corporativa soberana e inteligente para el modelado BPMN, estandarización Six Sigma (SIPOC) y gobierno de procesos operativos. Este Roadmap establece el cronograma maestro estructurado a lo largo del ciclo 2026 (16 de Enero de 2026 al 18 de Diciembre de 2026).")

    add_heading_1(doc, "2. Matriz de Hitos Clave del Proyecto TEMIS")
    tbl = doc.add_table(rows=8, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Fecha / Hito", "Fase Metodológica", "Entregable / Alcance Clave", "Estado"]
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    milestones = [
        ("16 Ene 2026: Planning Kickoff", "F1: Diagnóstico Estratégico", "Diagnóstico del dolor operativo en Lucidchart/Word y definición de arquitectura inicial.", "Completado"),
        ("15 Feb 2026: Project Charter", "F2: Inicio del Proyecto", "Acta Constitutiva, asignación de roles (PM, Sponsor) y delimitación de alcance.", "Completado"),
        ("31 Mar 2026: Backlog Scrum", "F3: Planificación Híbrida", "Backlog Scrum Técnico de 10 Sprints y arquitectura de experiencia.", "Completado"),
        ("01 Jun 2026: Migración Web SaaS", "F4: Ejecución Iterativa", "Hito Clave: Reestructuración mayor de App Windows Desktop a Web Cloud (Reflex + FastAPI + Postgres).", "Completado"),
        ("15 Sep 2026: Suite 4 Vistas & IA", "F4: Ejecución Iterativa", "Matriz SIPOC interactiva, generador de flujos Bézier, narrativa Gemini 2.5 Flash y catálogo en Archivo.", "Completado"),
        ("01 Nov 2026: Monitoreo & SA Drive", "F5/F6: Monitoreo y Mejora", "Auditor IA de calidad (0-100) y conexión con Service Account de Google Workspace Shared Drive.", "En curso"),
        ("18 Dic 2026: Cierre Oficial", "F7: Cierre del Proyecto", "Entrega formal de TEMIS v1.0, informe de lecciones aprendidas y traspaso operativo.", "Programado")
    ]

    for idx, m in enumerate(milestones, start=1):
        row = tbl.rows[idx]
        bg = "FFFFFF" if idx % 2 != 0 else HEX_LIGHT_BG
        for col_idx, text in enumerate(m):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            elif col_idx == 3:
                r.font.bold = True
                r.font.color.rgb = COLOR_SECONDARY if text == "Completado" else COLOR_PRIMARY

    add_heading_1(doc, "3. Desglose del Cronograma por las 7 Fases Corporativas")
    
    phases_info = [
        ("Fase 1: Diagnóstico Estratégico (Enero)", "Análisis del dolor en herramientas manuales desconectadas. Mapeo del proceso AS-IS de documentación de diagramas y definición de requerimientos de estandarización."),
        ("Fase 2: Inicio del Proyecto (Febrero)", "Publicación del Project Charter formal de TEMIS, aprobación de sponsor ejecutivo y establecimiento de la matriz RACI."),
        ("Fase 3: Planificación Híbrida (Marzo)", "Elaboración del Backlog Scrum Técnico 2026 con 47 commits base, especificación de simbología oficial BPMN y cálculo de esfuerzo."),
        ("Fase 4: Ejecución Iterativa (Abril - Septiembre)", "Desarrollo iterativo de la plataforma. Destaca la decisión estratégica de migrar de cliente Desktop Windows a arquitectura Web SaaS Full-Stack (Render Cloud), integración del motor SIPOC y Copiloto Gemini 2.5 Flash."),
        ("Fase 5: Monitoreo y Control (Octubre)", "Implementación del Auditor de Calidad IA (0-100) para cumplimiento de estándares Six Sigma, pruebas de carga y verificación continua."),
        ("Fase 6: Mejora Continua (Noviembre)", "Optimización de rendimiento en nube, refinamiento del catálogo de flujos guardados y automatización de respaldos en Google Workspace Shared Drive con Service Account."),
        ("Fase 7: Cierre del Proyecto (Diciembre)", "Evaluación de valor entregado, acta de cierre formal, paquete de exportación de proyectos y lecciones aprendidas.")
    ]
    for title_f, desc_f in phases_info:
        add_body_p(doc, desc_f, bold_prefix=f"• {title_f}: ")

    doc_path = os.path.join(OUTPUT_DIR, "Roadmap_Proyecto_TEMIS_2026.docx")
    doc.save(doc_path)
    print(f"✓ Guardado: {doc_path}")


# ==============================================================================
# 2. MARCO DE TRABAJO PARA LA GESTIÓN DEL PROYECTO TEMIS
# ==============================================================================
def create_framework_document():
    doc = docx.Document()
    add_document_header(
        doc,
        title="Marco de Trabajo para la Gestión del Proyecto TEMIS",
        subtitle="Metodología corporativa de gobernanza en 7 fases, matriz de roles, entregables y auditoría de calidad.",
        category="MARCO METODOLÓGICO & GOBERNANZA"
    )

    add_heading_1(doc, "1. Principios Rectores de TEMIS")
    add_body_p(doc, "1. Estandarización y Soberanía:", bold_prefix="• ")
    add_body_p(doc, "Eliminación de la dependencia de herramientas de terceros (Lucidchart, Miro, Visio) mediante un lienzo web soberano, determinístico y alineado a la simbología oficial BPMN.")
    add_body_p(doc, "2. Pipeline Automatizado de 3 Pasos:", bold_prefix="• ")
    add_body_p(doc, "Toda iniciativa se modela en una secuencia directa: Captura Tabular SIPOC ➔ Generación de Diagrama Bézier ➔ Redacción Automática de Narrativa / Manual con Gemini 2.5 Flash.")
    add_body_p(doc, "3. Gobernanza Continua y Calidad Six Sigma:", bold_prefix="• ")
    add_body_p(doc, "Auditoría estructural en tiempo real con calificación de 0 a 100 para garantizar que ningún proceso quede con decisiones inconclusas o actividades huérfanas.")

    add_heading_1(doc, "2. Estructura de Roles y Gobernanza")
    tbl = doc.add_table(rows=6, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Rol en el Proyecto", "Responsable / Área", "Responsabilidad Principal"]
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    roles = [
        ("Sponsor Ejecutivo", "Dirección de Operaciones & Tecnología", "Validación estratégica, asignación presupuestal y patrocinio corporativo."),
        ("Project Lead / PM", "Ing. Mario Hurtado", "Gestión integral del proyecto, cumplimiento del Roadmap 2026 y entregables."),
        ("Tech & Architecture Lead", "Equipo de Arquitectura de Software", "Diseño full-stack (Reflex, FastAPI, PostgreSQL) y rendimiento en Render."),
        ("AI & Automation Specialist", "Área de Innovación e IA", "Integración y prompts de Gemini 2.5 Flash para SIPOC, flujos y narrativa."),
        ("Process & QA Auditor", "Comité de Calidad y Procesos", "Revisión de manuales de procedimientos, auditoría de flujos y verificación UAT.")
    ]
    for idx, r_data in enumerate(roles, start=1):
        row = tbl.rows[idx]
        bg = "FFFFFF" if idx % 2 != 0 else HEX_LIGHT_BG
        for col_idx, text in enumerate(r_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY

    add_heading_1(doc, "3. El Ciclo de Vida Metodológico en 7 Fases")
    
    fases = [
        ("Fase 1: Diagnóstico Estratégico", "Entregables: Ficha de Diagnóstico, Matriz de Interesados y Justificación de Negocio. Foco: Identificación de ineficiencias en herramientas heredadas."),
        ("Fase 2: Inicio del Proyecto", "Entregables: Project Charter Oficial, Matriz SIPOC Inicial y Asignación RACI. Foco: Delimitación clara del alcance."),
        ("Fase 3: Planificación Híbrida", "Entregables: Diagrama BPMN Multi-Pestaña, Backlog Scrum Técnico 2026 y Matriz de Riesgos. Foco: Arquitectura y estimación."),
        ("Fase 4: Ejecución Iterativa", "Entregables: Manual de Políticas y Procedimientos, Servicios Backend/UI y Daily Logs. Foco: Sprints técnicos de desarrollo y despliegue continuo."),
        ("Fase 5: Monitoreo y Control", "Entregables: Auditoría IA de Calidad (0-100), Reporte de Cumplimiento SLA y Pruebas UAT. Foco: Semáforo del proyecto y gestión de bloqueos."),
        ("Fase 6: Mejora Continua", "Entregables: Plan de Ajustes Kaizen, Encuestas de Satisfacción y Métricas Operativas. Foco: Optimización continua de la suite."),
        ("Fase 7: Cierre del Proyecto", "Entregables: Acta de Cierre Aprobada, Paquete .temis.json Exportado y Lecciones Aprendidas. Foco: Traspaso a operaciones sostenibles.")
    ]
    for f_tit, f_desc in fases:
        add_body_p(doc, f_desc, bold_prefix=f"• {f_tit}: ")

    add_heading_1(doc, "4. Gestión de Riesgos y Control de Calidad")
    add_body_p(doc, "El marco de trabajo estipula la auditoría continua de procesos a través del motor inteligente de TEMIS, evaluando 4 dimensiones críticas: (1) Integridad de puntos de inicio/fin, (2) Completitud de compuertas de decisión Sí/No, (3) Asignación obligatoria de sistemas (Chronos/ERP) y canales (WhatsApp/Bria), y (4) Trazabilidad de entradas/salidas en la matriz SIPOC.")

    doc_path = os.path.join(OUTPUT_DIR, "Marco_de_Trabajo_Gestion_Proyecto_TEMIS.docx")
    doc.save(doc_path)
    print(f"✓ Guardado: {doc_path}")


# ==============================================================================
# 3. INFORME DIAGNÓSTICO GENERAL Y ANÁLISIS DE PROCESOS - PROYECTO TEMIS
# ==============================================================================
def create_diagnostic_document():
    doc = docx.Document()
    add_document_header(
        doc,
        title="Informe Diagnóstico General y Análisis de Procesos — Proyecto TEMIS",
        subtitle="Evaluación del estado actual (AS-IS) de documentación operativa vs la solución objetivo (TO-BE) de la Suite TEMIS.",
        category="INFORME DIAGNÓSTICO & ANÁLISIS DE PROCESOS"
    )

    add_heading_1(doc, "1. Introducción y Resumen Ejecutivo")
    add_body_p(doc, "La documentación y gobierno de procesos operativos ha dependido históricamente de herramientas desconectadas y procesos manuales que generan silos de información, desactualización de diagramas y horas de retrabajo en la redacción de manuales. El presente informe diagnostica las deficiencias del modelo actual (AS-IS) y fundamenta la arquitectura de solución de la Suite TEMIS (TO-BE).")

    add_heading_1(doc, "2. Diagnóstico del Estado Actual (AS-IS): Puntos de Dolor")
    
    pain_points = [
        ("Desconexión entre Diagramas y Manuales Escritos", "Un analista dibuja un flujo en Lucidchart o Visio y posteriormente tiene que redactar manualmente un documento de 20 páginas en Word, provocando inconsistencias severas entre lo dibujado y lo escrito."),
        ("Falta de Estandarización en Simbología y Carriles", "Uso heterogéneo de símbolos, falta de identificación clara de sistemas tecnológicos (ej. Chronos) y canales de comunicación (ej. WhatsApp)."),
        ("Captura Lenta y Compleja para Usuarios No Técnicos", "Diseñar un flujo directamente en un lienzo en blanco resulta intimidante y lento para analistas de negocio que no dominan herramientas CAD."),
        ("Ausencia de Auditoría de Calidad en Tiempo Real", "Los diagramas suelen contener compuertas de decisión sin salida alternativa o actividades huérfanas que no se detectan hasta una auditoría externa."),
        ("Limitación de Clientes de Escritorio Monousuario", "Las soluciones instaladas en Windows impiden la colaboración ágil, demandan instalaciones locales y dificultan la gestión centralizada de versiones.")
    ]
    for p_title, p_desc in pain_points:
        add_body_p(doc, p_desc, bold_prefix=f"• {p_title}: ")

    add_heading_1(doc, "3. Análisis Comparativo: Modelo Actual (AS-IS) vs Modelo TEMIS (TO-BE)")
    tbl = doc.add_table(rows=6, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Dimensión / Capacidad", "Modelo Tradicional (AS-IS)", "Modelo Suite TEMIS (TO-BE)"]
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    comparisons = [
        ("Entrada de Datos", "Dibujo manual caja por caja en lienzo en blanco.", "Matriz SIPOC Six Sigma tabular con autocompletado inteligente por Gemini 2.5 Flash."),
        ("Generación de Diagrama", "Horas de alineación manual de conectores y cajas.", "Generación automática en 1 clic con curvas Bézier perimetrales y etiquetas de sistemas/canales."),
        ("Documentación / Manual", "Redacción manual en Word propensa a errores.", "Redacción automática en prosa continua estructurada (1.0..N) con IA en segundos."),
        ("Auditoría y Calidad", "Revisión manual subjetiva y tardía.", "Auditor IA en tiempo real con puntuación Six Sigma 0-100 y recomendaciones automáticas."),
        ("Almacenamiento y Versiones", "Archivos dispersos en carpetas locales.", "Catálogo centralizado en Menú Archivo y respaldo automático en Google Workspace Shared Drive con SA.")
    ]
    for idx, c_data in enumerate(comparisons, start=1):
        row = tbl.rows[idx]
        bg = "FFFFFF" if idx % 2 != 0 else HEX_LIGHT_BG
        for col_idx, text in enumerate(c_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            elif col_idx == 2:
                r.font.bold = True
                r.font.color.rgb = COLOR_TEAL

    add_heading_1(doc, "4. Conclusiones y Plan de Transición")
    add_body_p(doc, "La implementación de TEMIS elimina el 80% del tiempo operativo dedicado a la diagramación y redacción de procedimientos, garantizando consistencia absoluta entre la captura tabular (SIPOC), la representación gráfica (BPMN) y el manual de políticas oficial. La plataforma se posiciona como el estándar corporativo definitivo para la gobernanza de procesos de la organización.")

    doc_path = os.path.join(OUTPUT_DIR, "Informe_Diagnostico_General_Proyecto_TEMIS.docx")
    doc.save(doc_path)
    print(f"✓ Guardado: {doc_path}")


if __name__ == "__main__":
    print("Generando los 3 documentos corporativos oficiales para PROYECTO TEMIS...")
    create_roadmap_document()
    create_framework_document()
    create_diagnostic_document()
    print("\n¡Los 3 documentos han sido generados exitosamente en Temis/docs/!")
